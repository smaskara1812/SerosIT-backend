"""The HAZ ID Register mail: every open Hazard ID card that has been open longer than the alert's Window (days), as an Excel file, mailed on the
alert's cycle (core/report_job.py). A port of the legacy alert 228 ("HAZ ID Register Mail Alert", Prc_Qry_Hazard_ID_Card_Report with
Get_Open_Haz_ID_Cards_Alert_Data and clsrHazard_Id_Cards_Dtl.Get_Open_Haz_ID_Cards_Alert_Data), with these differences, all agreed:
  * the age in days is the alert's Window field (the legacy procedure hard-coded 30; the alert's window is now 30), not a number in the code;
  * when nothing is open the mail still goes, saying so, without an attachment (the legacy code tried to attach a file it had not created);
  * a failure is retried and reported by e-mail to SCHEDULER_ERROR_MAIL_TO (the legacy code mailed the error text to the business-group admin);
  * the Excel is built here (openpyxl) with the same layout, columns and widths as the legacy file; the file name carries a 24-hour time (hhmmss, as in
    the legacy name); the wording says "Seros-IT" where the legacy file said "SIS" / "S.I.S.".
The cards: open ("O"), not deleted, event date at least `window` days before the run date; grouped by rig and contract; one Excel row per card."""

import datetime
import html
import io

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import HazardCard
from .report_job import ReportMail, ScheduledReportJob

JOB_KEY = "haz_id_register"
ENTITY_KEY = "automation"
ACTION = "send_haz_id_register_mail"
TRIGGER_CODE = "auto.haz_id_register"
ALERT_NAME = "HAZ ID Register Mail Alert"  # only what the setup command looks for when it is not told an alert id
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = [
    "SrNo", "Haz ID Card No", "Operator", "Location", "Date of Event", "Observer Name", "Hazard Type Name", "TOFS", "Hazard Description", "Action Taken",
    "Resp Dept", "Resp Rank", "Close Out Date", "Haz ID Status",
]
WIDTHS = {"SrNo": 5, "TOFS": 5, "Hazard Description": 30}  # every other column 15, as in the legacy file


def collect(window_days, today):
    """[(rig name, contract number, [row dict, ...]), ...] ordered by rig and contract, each group's cards by card number, SrNo restarting per group."""
    cutoff = today - datetime.timedelta(days=max(0, int(window_days or 0)))
    qs = (
        HazardCard.objects.filter(haz_id_card_status="O", event_dt__date__lte=cutoff).exclude(marked_as_deleted="Y")
        .select_related("rig", "contract", "contract__operator", "work_location", "haz_type", "resp_dept", "resp_rank")
        .order_by("rig_id", "contract_id", "haz_id_card_no", "haz_card_id")
    )
    groups, key = [], None
    for c in qs:
        if (c.rig_id, c.contract_id) != key:
            key = (c.rig_id, c.contract_id)
            groups.append((c.rig.rig_name, c.contract.prj_contract_no, []))
        rows = groups[-1][2]  # TOFS: the data keeps Y or blank, the legacy file says N
        rows.append({
            "SrNo": len(rows) + 1, "Haz ID Card No": c.haz_id_card_no, "Operator": c.contract.operator.operator_name, "Location": c.work_location.work_location,
            "Date of Event": f"{c.event_dt:%d/%m/%Y %H:%M}" if c.event_dt else "", "Observer Name": (c.reported_by_name or "").strip(),
            "Hazard Type Name": c.haz_type.haz_type_name, "TOFS": "Y" if c.timeout_for_safety == "Y" else "N", "Hazard Description": (c.hazard_desc or "").strip(),
            "Action Taken": (c.action_taken or "").strip(), "Resp Dept": c.resp_dept.vessel_dept_name, "Resp Rank": c.resp_rank.rank_name,
            "Close Out Date": f"{c.close_out_dt:%d/%m/%Y %H:%M}" if c.close_out_dt else "", "Haz ID Status": "Closed" if c.haz_id_card_status == "C" else "Open",
        })
    return groups


def build_workbook(groups, today):
    """The Excel bytes: a title, the report date, the header row, then for each rig a row with its name, a row with the contract number and its cards."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_properties.tabColor = "000000"
    ws.sheet_format.defaultRowHeight = 15
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    last = get_column_letter(len(HEADERS))  # N
    group_last = get_column_letter(len(HEADERS) - 1)  # M: the legacy file merges the rig and contract rows over the first 13 columns

    ws["A1"] = "Haz ID Card Register Report"
    ws["A1"].font = Font(name="Arial", size=15)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 18
    ws.merge_cells(f"A1:{last}1")
    ws["A2"] = CellRichText(TextBlock(InlineFont(b=True), "Seros-IT Report Date: "), f"{today:%d-%m-%Y}")
    ws["A2"].alignment = Alignment(horizontal="right")
    ws.merge_cells(f"A2:{last}2")
    ws["A3"].alignment = Alignment(horizontal="left")
    ws.merge_cells(f"A3:{last}3")

    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D3D3D3")
        c.alignment = Alignment(horizontal="center")
        c.border = box
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 15)

    r = 5
    for rig, contract, rows in groups:
        for text in (rig, contract):
            ws.cell(row=r, column=1, value=text)
            ws.merge_cells(f"A{r}:{group_last}{r}")
            r += 1
        for row in rows:
            for i, h in enumerate(HEADERS, start=1):
                c = ws.cell(row=r, column=i, value=row[h])
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = box
            r += 1
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _page(inner):
    return (
        "<div style=\"font-family:Arial,sans-serif;font-size:14px\">"
        "<p style=\"color:#e00707;margin:0 0 14px\">Seros-IT Auto Report</p>"
        f"{inner}"
        "<p style=\"margin:28px 0 0\">This is an automated Mail. Please do not reply.</p></div>"
    )


def build_mail(alert, now):
    """-> ReportMail for the alert (its Window gives the age in days, its Mail Subject the subject)."""
    window = max(0, int(alert.alert_window or 0))
    groups = collect(window, now.date())
    cards = sum(len(g[2]) for g in groups)
    subject = (alert.mail_subject or "Seros-IT Auto Report - HAZ ID Register").strip()
    if not cards:
        inner = f"<p>Dear Sir,</p><p>There are no open HAZ-ID cards that have been open for {window} days or more.</p>"
        return ReportMail(subject, _page(inner), "Dear Sir, there are no open HAZ-ID cards that have been open for %d days or more." % window, [], True, "Nothing open")
    name = f"Haz_Id_Register_{now:%d_%m_%Y_%H%M%S}.xlsx"
    inner = f"<p>Dear Sir,</p><p>Please find attached open HAZ-ID Register</p><p style=\"color:#666\">{cards} open cards, open {window} days or more, in {len(groups)} rig and contract groups.</p>"
    return ReportMail(
        subject, _page(inner), "Dear Sir, please find attached open HAZ-ID Register", [(name, build_workbook(groups, now.date()), XLSX)], False,
        f"{cards} open cards in {len(groups)} groups",
    )


class HazIdRegisterJob(ScheduledReportJob):
    key = JOB_KEY
    name = "HAZ ID Register"
    report_name = "HAZ ID Register"
    description = "Every Hazard ID card that is still open and was raised at least as many days ago as the alert's Window, as an Excel file grouped by rig and contract."
    trigger_entity = ENTITY_KEY
    trigger_action = ACTION
    trigger_code = TRIGGER_CODE

    def build(self, alert, now):
        return build_mail(alert, now)


JOB = HazIdRegisterJob()


def ensure_setup(alert_id=None, alert_name=ALERT_NAME):
    """Creates the scheduler's job row and, when the alert can be found (by id, else by name), the trigger row that connects the report to it. Safe to
    run again. Returns what it did, as a list of words."""
    from django.utils import timezone

    from .models import MailAlertDtl, NotificationTrigger, SchedJob

    did = []
    SchedJob.objects.get_or_create(job_key=JOB_KEY, defaults={"name": JOB.name, "enabled": "Y"})
    if NotificationTrigger.objects.filter(entity_key=ENTITY_KEY, action=ACTION).exists():
        return did
    alert = MailAlertDtl.objects.filter(pk=alert_id).first() if alert_id else MailAlertDtl.objects.filter(alert_name__iexact=alert_name).order_by("alert_id").first()
    if alert is None:
        return did
    NotificationTrigger.objects.create(entity_key=ENTITY_KEY, action=ACTION, alert=alert, notification_trigger_active="Y", cr_user_id=0, cr_dt=timezone.now())
    did.append(f"trigger row created for alert {alert.alert_id} ({alert.alert_name})")
    return did

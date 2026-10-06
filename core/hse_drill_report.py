"""QHSE → HSE Drill Report — rebuild of legacy's HSE Drill report form
(Prc_HSE_Emergency_Drill_Matrix). A form only; Print downloads an Excel
file with two parts: the Emergency Drill Matrix (one row per active drill,
weeks 1-52, Required, Total) and an "HSE Report" bar chart (drills x rigs).

Period -> date range, as in the legacy proc (calendar quarters/halves, not
financial year):
  MONTHLY     1st-last day of the picked month
  QUARTERLY   Jan-Mar / Apr-Jun / Jul-Sep / Oct-Dec of the picked year
  BIANNUAL    Jan-Jun / Jul-Dec of the picked year
  ANNUAL      Jan-Dec of the picked year
  DATE_RANGE  1st day of From month to last day of To month

Counts come from HSE Weekly Drill entries (HseWeeklyDrillDtl) by their
Drill Conducted Date, for the selected rigs. Matching legacy, a cell is
keyed by the header's week number only — NOT its year — so a Date Range
spanning two years adds week N of both years into the same column. Single-
period reports are unaffected. (Flagged for the mentor to decide on.)

"Required" is the number of times a year that drill's frequency calls for
(Weekly 52, Monthly 12, Quarterly 4, Bi-Annually 2, Annually 1) — a fixed
yearly figure, not scaled to the period, as in the legacy sample.

Letterhead: the general SEROS logo only, always, and no company name (the
legacy export has none).
"""

import calendar
import io
from collections import defaultdict
from datetime import date, datetime

from django.db.models import Count
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework.response import Response
from rest_framework.views import APIView

from . import audit as _audit
from .company_branding import seros_logo_path
from .models import HseWeeklyDrillDtl, MstHseDrill, MstRig
from .permissions import HasMenuPermission

ENTITY_KEY = "qhse.hse_drill_report"
REQUIRED_PER_YEAR = {"W": 52, "M": 12, "Q": 4, "B": 2, "A": 1}
QUARTERS = {"JAN-MAR": 0, "APR-JUN": 3, "JUL-SEP": 6, "OCT-DEC": 9}
HALVES = {"JAN-JUN": 0, "JUL-DEC": 6}
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _month_start(year, month):
    return date(year, month, 1)


def _month_end(year, month):
    return date(year, month, calendar.monthrange(year, month)[1])


def _add_months(year, month, n):
    idx = year * 12 + (month - 1) + n
    return idx // 12, idx % 12 + 1


def _parse_month(value, label):
    try:
        y, m = value.split("-")
        y, m = int(y), int(m)
        if not (1 <= m <= 12 and 1900 <= y <= 2100):
            raise ValueError
        return y, m
    except (AttributeError, ValueError):
        raise ValueError(f"{label} must be a valid month and year.")


def _parse_year(value):
    try:
        y = int(value)
        if not 1900 <= y <= 2100:
            raise ValueError
        return y
    except (TypeError, ValueError):
        raise ValueError("Year must be a valid 4-digit year.")


def resolve_period(params):
    """-> (from_date, to_date, label). Raises ValueError with a user-facing message."""
    period = params.get("period")
    if period == "MONTHLY":
        y, m = _parse_month(params.get("month"), "Month/Year")
        return _month_start(y, m), _month_end(y, m), f"{MON[m - 1]}-{y}"
    if period in ("QUARTERLY", "BIANNUAL"):
        y = _parse_year(params.get("year"))
        table, span, name = (QUARTERS, 2, "Quarter") if period == "QUARTERLY" else (HALVES, 5, "Months")
        sub = params.get("sub")
        if sub not in table:
            raise ValueError(f"{name} must be selected.")
        sy, sm = _add_months(y, 1, table[sub])
        ey, em = _add_months(sy, sm, span)
        return _month_start(sy, sm), _month_end(ey, em), f"{MON[sm - 1]}-{sy} To {MON[em - 1]}-{ey}"
    if period == "ANNUAL":
        y = _parse_year(params.get("year"))
        return date(y, 1, 1), date(y, 12, 31), f"Jan-{y} To Dec-{y}"
    if period == "DATE_RANGE":
        fy, fm = _parse_month(params.get("month"), "Month/Year")
        ty, tm = _parse_month(params.get("to_month"), "To Month/Year")
        start, end = _month_start(fy, fm), _month_end(ty, tm)
        if start > end:
            raise ValueError("To Month/Year can't be before Month/Year.")
        return start, end, f"{MON[fm - 1]}-{fy} To {MON[tm - 1]}-{ty}"
    raise ValueError("Period must be selected.")


def build_report(params):
    d1, d2, label = resolve_period(params)
    try:
        rig_ids = sorted({int(x) for x in (params.get("rigs") or "").split(",") if x.strip()})
    except ValueError:
        raise ValueError("Invalid rig selection.")
    if not rig_ids:
        raise ValueError("Select at least one rig.")
    rigs = list(MstRig.objects.filter(pk__in=rig_ids).order_by("rig_name"))
    if not rigs:
        raise ValueError("Invalid rig selection.")

    drills = list(MstHseDrill.objects.filter(hse_drill_active="Y").order_by("hse_drill_name", "hse_drill_id"))
    base = HseWeeklyDrillDtl.objects.filter(drill_conducted_dt__range=(d1, d2), hdr__rig_id__in=[r.pk for r in rigs])

    by_week = defaultdict(int)
    for row in base.values("hse_drill_id", "hdr__drill_week").annotate(n=Count("pk")):
        by_week[(row["hse_drill_id"], row["hdr__drill_week"])] = row["n"]
    matrix = []
    for d in drills:
        weeks = [by_week.get((d.pk, w)) or None for w in range(1, 53)]
        matrix.append(
            {
                "name": d.hse_drill_name,
                "frequency": d.get_hse_drill_frequency_display(),
                "weeks": weeks,
                "required": REQUIRED_PER_YEAR.get(d.hse_drill_frequency),
                "total": sum(w or 0 for w in weeks) or None,
            }
        )

    per_rig = defaultdict(int)
    for row in base.values("hdr__rig__rig_name", "hse_drill__hse_drill_name").annotate(n=Count("pk")):
        per_rig[(row["hse_drill__hse_drill_name"], row["hdr__rig__rig_name"])] += row["n"]
    chart_drills = sorted({k[0] for k in per_rig})
    chart_rigs = sorted({k[1] for k in per_rig})
    return {
        "from": d1, "to": d2, "label": label, "rigs": rigs, "matrix": matrix,
        "chart": {"drills": chart_drills, "rigs": chart_rigs, "counts": per_rig},
    }


def render_workbook(report):
    """Laid out to match the legacy export: logo top-left, "SIS Report Date"
    top-right, a rule under row 4, the bold title in row 6, a grey header row
    (yellow Required/Total), tight bordered rows, then the chart straight
    after its title. The chart's source numbers live on a hidden second
    sheet so only the chart shows, as in legacy."""
    wb = Workbook()
    ws = wb.active
    ws.title = "HSE_Emergency_Drill_Matrix"
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    yellow = PatternFill("solid", fgColor="FFFF00")
    base_font = Font(name="Arial", size=9)
    bold_font = Font(name="Arial", size=9, bold=True)

    first_col, last_col = 3, 3 + 1 + 52 + 2  # C .. Total
    # Only the general SEROS logo — never a per-company one.
    logo_abs = seros_logo_path()
    if logo_abs:
        img = XlImage(logo_abs)
        img.height, img.width = 36, int(36 * img.width / img.height)
        ws.add_image(img, "C1")
    ws.merge_cells(start_row=2, start_column=last_col - 14, end_row=2, end_column=last_col)
    c = ws.cell(row=2, column=last_col - 14, value=f"SIS Report Date:  {datetime.now():%d/%m/%Y %H:%M}")
    c.font, c.alignment = Font(name="Arial", size=8, bold=True), Alignment(horizontal="right")
    for col in range(first_col, last_col + 1):
        ws.cell(row=4, column=col).border = Border(bottom=Side(style="medium", color="808080"))

    ws.cell(row=6, column=first_col, value=f"Emergency Drill Matrix ({report['label']} ) ({report['from']:%d/%m/%Y} - {report['to']:%d/%m/%Y})").font = Font(
        name="Arial", size=10, bold=True
    )

    hdr_row = 8
    headers = ["Drill Name", "Drill Frequency", *range(1, 53), "Required", "Total"]
    for i, h in enumerate(headers):
        c = ws.cell(row=hdr_row, column=first_col + i, value=h)
        c.font, c.border = bold_font, border
        c.alignment = Alignment(horizontal="left" if i < 2 else "center", vertical="center")
        c.fill = yellow if h in ("Required", "Total") else head_fill
    for r, m in enumerate(report["matrix"], start=hdr_row + 1):
        vals = [m["name"], m["frequency"], *m["weeks"], m["required"], m["total"]]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=first_col + i, value=v)
            c.font, c.border = base_font, border
            c.alignment = Alignment(horizontal="left" if i < 2 else "center")
        ws.row_dimensions[r].height = 12.75
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 2
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 14
    for col in range(first_col + 2, first_col + 2 + 52):
        ws.column_dimensions[get_column_letter(col)].width = 3.6
    ws.column_dimensions[get_column_letter(last_col - 1)].width = 9
    ws.column_dimensions[get_column_letter(last_col)].width = 6.5

    chart = report["chart"]
    title_row = hdr_row + len(report["matrix"]) + 3
    ws.cell(row=title_row, column=first_col, value=f"HSE Report ({report['label']} ) ({report['from']:%d/%m/%Y} - {report['to']:%d/%m/%Y})").font = Font(
        name="Arial", size=10, bold=True
    )
    if chart["drills"]:
        data = wb.create_sheet("Chart Data")
        data.cell(row=1, column=1, value="Drill")
        for j, rig in enumerate(chart["rigs"]):
            data.cell(row=1, column=2 + j, value=rig)
        for i, drill in enumerate(chart["drills"]):
            data.cell(row=2 + i, column=1, value=drill)
            for j, rig in enumerate(chart["rigs"]):
                data.cell(row=2 + i, column=2 + j, value=chart["counts"].get((drill, rig), 0))
        data.sheet_state = "hidden"
        bar = BarChart()
        bar.type, bar.grouping = "col", "clustered"
        bar.title = "HSE Report"
        bar.x_axis.title, bar.y_axis.title = "Emergency Drill", "No. of drills"
        last = 1 + len(chart["drills"])
        bar.add_data(Reference(data, min_col=2, max_col=1 + len(chart["rigs"]), min_row=1, max_row=last), titles_from_data=True)
        bar.set_categories(Reference(data, min_col=1, min_row=2, max_row=last))
        bar.height, bar.width = 9, 22
        bar.x_axis.delete = bar.y_axis.delete = False
        ws.add_chart(bar, f"C{title_row + 1}")
    else:
        ws.cell(row=title_row + 1, column=first_col, value="No drills conducted in this period for the selected rigs.").font = base_font

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class HseDrillReportExportView(APIView):
    permission_classes = [HasMenuPermission]
    entity_key = ENTITY_KEY
    # Same reasoning as IncidentDashboardExportView: a plain APIView has no
    # router-assigned .action, so set it to require the Export permission.
    action = "export"

    def get(self, request):
        try:
            report = build_report(request.query_params)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=400)
        content = render_workbook(report)
        response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="HSE Emergency Drill Matrix.xlsx"'
        _audit.record_action(
            request, "export", ENTITY_KEY,
            record_label=f"HSE Drill Report — {report['label']}",
            changes={
                "period": {"old": None, "new": request.query_params.get("period")},
                "from": {"old": None, "new": report["from"].isoformat()},
                "to": {"old": None, "new": report["to"].isoformat()},
                "rigs": {"old": None, "new": ", ".join(r.rig_name for r in report["rigs"])},
            },
        )
        return response

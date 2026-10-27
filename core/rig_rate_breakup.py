"""The Rig Rate Break-Up Sheet mail: one Excel file with one sheet per active rig, month to date, day by day, hours per rate. A port of the legacy
alert 233 ("Rig Rate Break-Up Sheet": Prc_Qry_Rig_Revenue_Report with Get_Rig_Revenue_All_Rigs_Data, and Get_Rig_Revenue_Details / get_Tables that
wrote the Excel), mailed on the alert's cycle (core/report_job.py): on the alert's Next Mail date, at its Start Time.

What the sheets contain (the same as the legacy report):
  * Dates: the 1st of the month up to the day it is sent. On the 1st of a month, the whole previous month.
  * Rigs: a rig is in it when it has an active contract line (no end date on the line or on the contract; the latest line of the rig), its cost
    centre is active, and it has a daily drilling report on the contract between (yesterday minus the alert's Window days) and yesterday.
    The legacy procedure hard-coded those 30 days; here it is the alert's Window field.
  * One row per date: Date, Well (the drilling record's location on that date), the hours per rate R1, R2, R3, R0, FM, ILM, IPM - Lateral
    (drilling rates 1 to 7, summed from the day's activity blocks), Total, and the Operations Summary (the day's remark). A day with no report is a
    blank row, as in the legacy file (today's row is included). A last row, "Total Hours", adds each column up.
  * The header: the rig's company on that date, "Rig Rate Break-Up Sheet", the report date, the rig, its locations in the period, its area, the month.
Differences from the legacy file, all agreed: the sheets are named after their rig (not Sheet1, Sheet2); the wording says "Seros-IT" where the legacy
file said "SIS" / "S.I.S."; when no rig qualifies the mail still goes, saying so, with no attachment."""

import calendar
import datetime
import io
import re
from collections import defaultdict
from decimal import Decimal

from django.db.models import Sum
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .company_branding import resolve_rig_company_branding
from .models import DrillingDtl, DrillingDtlOps, DrillingHdr, MstCostCentre, MstDrillingRate, MstRig, ProjectContractDtl
from .report_job import ReportMail, ScheduledReportJob

JOB_KEY = "rig_rate_breakup"
ENTITY_KEY = "automation"
ACTION = "send_rig_rate_breakup_mail"
TRIGGER_CODE = "auto.rig_rate_breakup"
ALERT_NAME = "Rig Rate Break-Up Sheet"  # only what the setup command looks for when it is not told an alert id
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RATE_IDS = (1, 2, 3, 4, 5, 6, 7)  # the drilling rates the sheet breaks down (R1, R2, R3, R0, FM, ILM, IPM - Lateral), as in the legacy procedure
OPS_DRILL_ACTUAL = 2  # Mst_Drilling_Operation: "Drill actual"
METERAGE_RATES = {13: 3, 14: 4}  # rate id -> its drilling section (12-1/4", 8-1/2"): a day that has only these still shows its remark
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def period(today):
    """(from, to) of the sheet: the 1st of this month to `today`; on the 1st, the whole previous month."""
    if today.day == 1:
        last = today - datetime.timedelta(days=1)
        return last.replace(day=1), last
    return today.replace(day=1), today


def month_label(d):
    return f"{MONTHS[d.month - 1]}-{d.year}"


def active_rigs(today, lookback_days):
    """The rigs that get a sheet, ordered by rig id: [MstRig, ...]."""
    yesterday = today - datetime.timedelta(days=1)
    back = yesterday - datetime.timedelta(days=max(0, int(lookback_days or 0)))
    latest = {}
    for ln in ProjectContractDtl.objects.filter(rig_active_to__isnull=True, contract__prj_end_dt__isnull=True):
        cur = latest.get(ln.rig_id)
        if cur is None or (ln.rig_active_from, ln.pk) > (cur.rig_active_from, cur.pk):
            latest[ln.rig_id] = ln
    out = []
    for rig_id, ln in latest.items():
        if not MstCostCentre.objects.filter(rig_id=rig_id, cost_centre_active="Y").exists():
            continue
        if DrillingDtl.objects.filter(rig_id=rig_id, drilling_hdr__contract_id=ln.contract_id, drilling_dtl_dt__gte=back, drilling_dtl_dt__lte=yesterday).exists():
            out.append(rig_id)
    return list(MstRig.objects.filter(pk__in=out).order_by("rig_id"))


def rate_codes():
    return {r.drilling_rate_id: (r.rate_code or "").strip() for r in MstDrillingRate.objects.filter(drilling_rate_id__in=RATE_IDS)}


def rig_sheet(rig, start, end):
    """Everything one sheet shows: {"company", "rig", "locations", "area", "month", "rows": [...], "totals": {...}}.
    Each row: {"date", "well", "rates": {id: Decimal or None}, "total": Decimal or None, "remark"}; a day with no report has rates/total None."""
    hdrs = list(DrillingHdr.objects.filter(rig_id=rig.rig_id).select_related("contract", "contract__location"))

    def anchor(h):
        return h.first_anchor_down_dt.date()

    def done(h):
        return h.drilling_completion_dt.date() if h.drilling_completion_dt else end

    in_period = [h for h in hdrs if done(h) >= start and anchor(h) <= end]
    in_period.sort(key=lambda h: (anchor(h), h.pk))
    locations = ", ".join(dict.fromkeys(h.location for h in in_period if h.location))
    area = in_period[-1].contract.location.location_name if in_period and in_period[-1].contract_id else ""

    remarks = {d.drilling_dtl_dt: (d.remark or "").strip() for d in DrillingDtl.objects.filter(rig_id=rig.rig_id, drilling_dtl_dt__gte=start, drilling_dtl_dt__lte=end)}
    hours = defaultdict(lambda: defaultdict(Decimal))
    has_row = set()
    ops = (
        DrillingDtlOps.objects.filter(drilling_dtl__rig_id=rig.rig_id, drilling_dtl__drilling_dtl_dt__gte=start, drilling_dtl__drilling_dtl_dt__lte=end)
        .values("drilling_dtl__drilling_dtl_dt", "prj_drilling_rate__drilling_rate_id", "drilling_ops_id", "drilling_section_id")
        .annotate(h=Sum("duration"))
    )
    for o in ops:
        day, rid = o["drilling_dtl__drilling_dtl_dt"], o["prj_drilling_rate__drilling_rate_id"]
        if rid in RATE_IDS:
            hours[day][rid] += o["h"] or Decimal(0)
            has_row.add(day)
        elif rid in METERAGE_RATES and o["drilling_ops_id"] == OPS_DRILL_ACTUAL and o["drilling_section_id"] == METERAGE_RATES[rid]:
            has_row.add(day)  # the day has a report (metres drilled), but no hours in the seven columns: its remark and a total of 0 are shown

    rows, totals = [], {rid: Decimal(0) for rid in RATE_IDS}
    grand = Decimal(0)
    day = start
    while day <= end:
        wells = [h for h in hdrs if anchor(h) <= day <= done(h)]
        well = max(wells, key=lambda h: (anchor(h), h.pk)).location if wells else ""
        if day in has_row:
            per = {rid: (hours[day][rid] if hours[day].get(rid) else None) for rid in RATE_IDS}
            total = sum((v for v in per.values() if v), Decimal(0))
            for rid, v in per.items():
                if v:
                    totals[rid] += v
            grand += total
            rows.append({"date": day, "well": well, "rates": per, "total": total, "remark": remarks.get(day, "")})
        else:
            rows.append({"date": day, "well": well, "rates": {rid: None for rid in RATE_IDS}, "total": None, "remark": ""})
        day += datetime.timedelta(days=1)
    branding = resolve_rig_company_branding(rig.rig_id, end)
    return {
        "company": branding.get("company_name") or "", "rig": rig.rig_name, "locations": locations, "area": area, "month": month_label(end),
        "rows": rows, "totals": {rid: (v or None) for rid, v in totals.items()}, "grand": grand or None,
    }


def collect(today, lookback_days):
    """[sheet dict, ...] for every rig that qualifies on `today`, and the period (from, to)."""
    start, end = period(today)
    return [rig_sheet(r, start, end) for r in active_rigs(today, lookback_days)], (start, end)


def _num(v):
    return None if v is None else float(v)


def _sheet_name(name, taken):
    base = re.sub(r"[\[\]:*?/\\]", " ", name or "Rig").strip()[:28] or "Rig"
    out, n = base, 2
    while out.lower() in taken:
        out = f"{base} ({n})"
        n += 1
    taken.add(out.lower())
    return out


def build_workbook(sheets, today, codes=None):
    """The Excel bytes: one sheet per rig, laid out like the legacy file."""
    codes = codes or rate_codes()
    heads = ["Date", "Well"] + [codes.get(i, f"Rate {i}") for i in RATE_IDS] + ["Total", "Operations Summary"]
    ncol = len(heads)  # 11
    last = get_column_letter(ncol)
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    wb.remove(wb.active)
    taken = set()
    for sh in sheets:
        ws = wb.create_sheet(_sheet_name(sh["rig"], taken))
        ws.sheet_format.defaultRowHeight = 15
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws["A1"] = sh["company"]
        ws["A1"].font = Font(name="Calibri", size=15, bold=True, color="0000FF")
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A2"] = "Rig Rate Break-Up Sheet"
        ws["A2"].font = Font(name="Calibri", size=14, bold=True)
        ws["A2"].alignment = Alignment(horizontal="center")
        ws["A3"] = CellRichText(TextBlock(InlineFont(b=True), "Seros-IT Report Date: "), f"{today:%d-%m-%Y}")
        ws["A4"] = CellRichText(TextBlock(InlineFont(b=True), "Rig:  "), sh["rig"])
        ws["A5"] = CellRichText(TextBlock(InlineFont(b=True), "Locations : "), sh["locations"])
        ws["A6"] = CellRichText(TextBlock(InlineFont(b=True), "Area: "), sh["area"])
        for r in (3, 4, 5, 6):
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="right")
        ws["A7"] = sh["month"]
        ws["A7"].font = Font(name="Calibri", size=13, bold=True)
        ws["A7"].alignment = Alignment(horizontal="center")
        for r in range(1, 10):
            ws.merge_cells(f"A{r}:{last}{r}")
        for i, h in enumerate(heads, start=1):
            c = ws.cell(row=10, column=i, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D3D3D3")
            c.alignment = Alignment(horizontal="center")
            c.border = box
            ws.column_dimensions[get_column_letter(i)].width = 45 if h == "Operations Summary" else 17 if h == "IPM - Lateral" else 15
        r = 11
        for row in sh["rows"]:
            vals = [f"{row['date']:%d/%m/%Y}", row["well"]] + [_num(row["rates"][i]) for i in RATE_IDS] + [_num(row["total"]), row["remark"]]
            r = _put(ws, r, vals, box)
        _put(ws, r, ["Total Hours", ""] + [_num(sh["totals"][i]) for i in RATE_IDS] + [_num(sh["grand"]), ""], box)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _put(ws, r, vals, box):
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = box
    return r + 1


def _page(inner):
    return (
        "<div style=\"font-family:Arial,sans-serif;font-size:14px\">"
        "<p style=\"color:#e00707;margin:0 0 14px\">Seros-IT Auto Report</p>"
        f"{inner}"
        "<p style=\"margin:28px 0 0\">This is an automated Mail. Please do not reply.</p></div>"
    )


def build_mail(alert, now):
    """-> ReportMail: the workbook with a sheet per rig (the alert's Window gives how far back a rig must have a daily report)."""
    window = max(0, int(alert.alert_window or 0))
    sheets, (start, end) = collect(now.date(), window)
    subject = (alert.mail_subject or "Seros-IT Auto Report - Rig Rate Break-Up Sheet").strip()
    if not sheets:
        inner = f"<p>Dear Sir,</p><p>No rig has a daily drilling report in the last {window} days, so there is no Rig Rate Break-Up Sheet to send.</p>"
        return ReportMail(subject, _page(inner), f"Dear Sir, no rig has a daily drilling report in the last {window} days.", [], True, "No rig with a daily report")
    name = f"Rig_Rate_Break_Up_Sheet_{now:%d_%m_%Y_%H%M%S}.xlsx"
    inner = f"<p>Dear Sir,</p><p>Please find attached Rig Rate Break-Up Sheet dated {now:%d-%m-%Y}</p><p style=\"color:#666\">{len(sheets)} rigs, {start:%d %b} to {end:%d %b %Y}.</p>"
    return ReportMail(
        subject, _page(inner), f"Dear Sir, please find attached Rig Rate Break-Up Sheet dated {now:%d-%m-%Y}", [(name, build_workbook(sheets, now.date()), XLSX)], False,
        f"{len(sheets)} rigs, {start:%d %b} to {end:%d %b %Y}",
    )


class RigRateBreakUpJob(ScheduledReportJob):
    key = JOB_KEY
    name = "Rig Rate Break-Up"
    report_name = "Rig Rate Break-Up Sheet"
    description = "One Excel file with a sheet per active rig: the hours per rate for each day of the month so far (the whole previous month when it is sent on the 1st)."
    trigger_entity = ENTITY_KEY
    trigger_action = ACTION
    trigger_code = TRIGGER_CODE

    def build(self, alert, now):
        return build_mail(alert, now)

    def get_setup(self, today=None):
        s = super().get_setup(today)
        if s["alert"] and not (s["window"] or 0) >= 1:
            s["warnings"].append("The alert's Window is 0: only rigs with a daily report yesterday are included. The legacy report used 30 days; set the Window on the alert.")
        return s


JOB = RigRateBreakUpJob()


def ensure_setup(alert_id=None, alert_name=ALERT_NAME):
    """Creates the scheduler's job row and, when the alert can be found (by id, else by name), the trigger row that connects the report to it. Safe to
    run again. Returns what it did, as a list of words."""
    from django.utils import timezone

    from .models import MailAlertDtl, NotificationTrigger, SchedJob

    did = []
    SchedJob.objects.get_or_create(job_key=JOB_KEY, defaults={"name": JOB.name, "enabled": "Y"})
    if NotificationTrigger.objects.filter(entity_key=ENTITY_KEY, action=ACTION).exists():
        return did
    alert = MailAlertDtl.objects.filter(pk=alert_id).first() if alert_id else MailAlertDtl.objects.filter(alert_name__iexact=alert_name).order_by("alert_id").first()
    if alert is None:
        return did
    NotificationTrigger.objects.create(entity_key=ENTITY_KEY, action=ACTION, alert=alert, notification_trigger_active="Y", cr_user_id=0, cr_dt=timezone.now())
    did.append(f"trigger row created for alert {alert.alert_id} ({alert.alert_name})")
    return did

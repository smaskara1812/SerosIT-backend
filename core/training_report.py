"""QHSE → Training Report — rebuild of legacy rfrmHSE_Training_Report. A form
only; Print downloads an Excel file. Two report types:

  HSE Training Matrix (Prc_HSE_Training_Matrix)
    Designations (ranks with an active certificate mapping in the chosen
    category) down the side, certificates across the top. A cell is "M" when
    that certificate applies to the rank and some Training Group marks the
    rank's training mandatory. Certificates: Internal first, then External,
    then by name; external ones get a grey heading.

  Employees's HSE Training Matrix (Prc_Qry_Employee_HSE_Matrix)
    Needs Rig, From and To Date (To can't be in the future, From can't be after
    To). One row per FS employee of the category on that rig whose leaving
    date is blank or inside the range, in rank order. A cell is "M" when the
    certificate applies to the employee's rank, or the training date(s) when
    they attended that course at that rig in the range. Certificates follow the
    legacy order (type, then id).

Differences from legacy, agreed with the user:
  - One row per employee. Legacy repeated a person's row once per training they
    attended; here several dates share the cell, comma-separated.
  - The summary lines are named for what they count: "Candidates requiring
    training" (the M cells — legacy called this "Total Candidates Trained"),
    "Candidates trained" (cells with a date) and "Sessions held" (the real count
    of training logs; legacy hard-coded 0).
  - A real .xlsx with the general SEROS logo, no company name.
Categories offered follow the user's category mapping; admins see all.
"""

import io
from datetime import date, datetime

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework.response import Response
from rest_framework.views import APIView

from . import audit as _audit
from .company_branding import seros_logo_path
from .models import (
    CertToRankMapping,
    FsCatgToRankMapping,
    HseTrainingGroupDtl,
    HseTrainingLogDtl,
    HseTrainingLogHdr,
    MstCert,
    MstFsCategory,
    MstFsEmployee,
    MstRank,
    MstRig,
)
from .permissions import HasMenuPermission
from .training_group import allowed_category_ids

ENTITY_KEY = "qhse.training_report"
TYPES = {"Training_Matrix": "HSE Training Matrix", "Employees_Matrix": "Employee HSE Matrix"}
_TYPE_RANK = {"I": 0, "E": 1}  # Internal before External before untyped


def _cert_sort(cert, by):
    return (_TYPE_RANK.get(cert.cert_training_type, 2), cert.cert_name if by == "name" else cert.pk)


def _certs_for(category_id, by, oilfield_only):
    ids = CertToRankMapping.objects.filter(fs_category_id=category_id, cert_to_rank_mapping_active="Y").values_list("cert_id", flat=True)
    qs = MstCert.objects.filter(pk__in=set(ids))
    if oilfield_only:
        qs = qs.filter(business_system_id_6="Y")
    return sorted(qs, key=lambda c: _cert_sort(c, by))


def _category_label(category):
    return category.fs_category_name.upper().replace(" RIG PERSONNEL", "")


def _parse_date(value, label):
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label} must be a valid date.")


def build_report(request):
    p = request.query_params
    report_type = p.get("report_type")
    if report_type not in TYPES:
        raise ValueError("Report Type must be selected.")
    category_id = p.get("category")
    if not (category_id and category_id.isdigit()):
        raise ValueError("Category must be selected.")
    allowed = allowed_category_ids(request)
    category = MstFsCategory.objects.filter(pk=category_id).first()
    if not category or (allowed is not None and category.pk not in allowed):
        raise ValueError("Category must be selected.")
    if report_type == "Training_Matrix":
        return {"type": report_type, "category": category, **_matrix(category)}

    errors = []
    rig = MstRig.objects.filter(pk=p["rig"]).first() if (p.get("rig") or "").isdigit() else None
    if not rig:
        errors.append("Rig must be selected.")
    start = end = None
    try:
        start = _parse_date(p.get("from_date"), "From Date")
    except ValueError as exc:
        errors.append("From Date must be entered." if not p.get("from_date") else str(exc))
    try:
        end = _parse_date(p.get("to_date"), "To Date")
        if end > timezone.now().date():
            errors.append("To Date can't be in the future.")
    except ValueError as exc:
        errors.append("To Date must be entered." if not p.get("to_date") else str(exc))
    if start and end and end < start:
        errors.append("To Date can't be before From Date.")
    if errors:
        raise ValueError("\n".join(errors))
    return {"type": report_type, "category": category, "rig": rig, "from": start, "to": end, **_employees(category, rig, start, end)}


def _matrix(category):
    certs = _certs_for(category.pk, "name", oilfield_only=True)
    mapped = set(
        CertToRankMapping.objects.filter(fs_category=category, cert_to_rank_mapping_active="Y").values_list("cert_id", "rank_id")
    )
    mandatory = set(HseTrainingGroupDtl.objects.filter(fs_category=category, mandatory_training="Y").values_list("rank_id", flat=True))
    rank_ids = {r for _, r in mapped}
    ranks = sorted(MstRank.objects.filter(pk__in=rank_ids), key=lambda r: r.rank_name)
    rows = [(r.rank_name, [("M" if (c.pk, r.pk) in mapped and r.pk in mandatory else "") for c in certs]) for r in ranks]
    return {"certs": certs, "rows": rows}


def _employees(category, rig, start, end):
    certs = _certs_for(category.pk, "id", oilfield_only=False)
    cert_ids = [c.pk for c in certs]
    mapped = set(
        CertToRankMapping.objects.filter(fs_category=category, cert_to_rank_mapping_active="Y").values_list("cert_id", "rank_id")
    )
    order = {}
    for m in FsCatgToRankMapping.objects.filter(fs_category=category).values("rank_id", "rank_order"):
        order[m["rank_id"]] = min(order.get(m["rank_id"], 10**6), m["rank_order"] if m["rank_order"] is not None else 10**6)

    staff = list(
        MstFsEmployee.objects.filter(fs_category=category, rig=rig)
        # A blank leaving date counts as still here (legacy used the To Date for it).
        .filter(Q(fs_emp_dol__isnull=True) | Q(fs_emp_dol__range=(start, end)))
        .select_related("rank")
    )
    staff.sort(key=lambda e: (order.get(e.rank_id, 10**6), e.rank.rank_name, e.fs_emp_lname, e.fs_emp_fname or ""))

    trained = {}
    logs = HseTrainingLogDtl.objects.filter(
        hdr__rig=rig, hdr__training_dt__range=(start, end), hdr__cert_id__in=cert_ids, fs_emp_id__in=[e.pk for e in staff]
    ).values_list("fs_emp_id", "hdr__cert_id", "hdr__training_dt")
    for emp_id, cert_id, when in logs:
        trained.setdefault((emp_id, cert_id), []).append(when)

    rows, requiring, done = [], [0] * len(certs), [0] * len(certs)
    for i, e in enumerate(staff, start=1):
        cells = []
        for j, c in enumerate(certs):
            dates = sorted(trained.get((e.pk, c.pk), []))
            if dates:
                cells.append(", ".join(d.strftime("%d/%m/%Y") for d in dates))
                done[j] += 1
            elif (c.pk, e.rank_id) in mapped:
                cells.append("M")
                requiring[j] += 1
            else:
                cells.append("")
        rows.append((i, str(e.fs_emp_staff_id or ""), str(e), e.rank.rank_name, cells))
    sessions = [
        HseTrainingLogHdr.objects.filter(rig=rig, cert_id=c.pk, training_dt__range=(start, end)).count() for c in certs
    ]
    return {"certs": certs, "rows": rows, "requiring": requiring, "trained": done, "sessions": sessions}


# ── Excel ─────────────────────────────────────────────────────────────────
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GREY = PatternFill("solid", fgColor="D3D3D3")
_BASE = Font(name="Arial", size=9)
_BOLD = Font(name="Arial", size=9, bold=True)


def _logo(ws, anchor):
    path = seros_logo_path()
    if path:
        img = XlImage(path)
        img.height, img.width = 40, int(40 * img.width / img.height)
        ws.add_image(img, anchor)


def _render_matrix(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "HSE_Training_Matrix"
    certs, last_col = report["certs"], 2 + len(report["certs"])
    _logo(ws, "A1")
    title = ws.cell(row=2, column=max(3, last_col), value=f"HSE TRAINING MATRIX ({_category_label(report['category'])})")
    title.font, title.alignment = Font(name="Arial", size=11, bold=True), Alignment(horizontal="right")
    for col in range(1, last_col + 1):
        ws.cell(row=4, column=col).border = Border(bottom=Side(style="medium", color="808080"))
    head = 6
    ws.row_dimensions[head].height = 170
    for col, text in ((1, "Sr. No."), (2, "Designation / Certificates")):
        c = ws.cell(row=head, column=col, value=text)
        c.font, c.border, c.alignment = _BOLD, _BORDER, Alignment(vertical="top", wrap_text=True)
    for j, cert in enumerate(certs, start=3):
        c = ws.cell(row=head, column=j, value=cert.cert_name)
        c.font, c.border = _BOLD, _BORDER
        c.alignment = Alignment(textRotation=90, horizontal="center", vertical="bottom")
        if cert.cert_training_type == "E":
            c.fill = _GREY
        ws.column_dimensions[get_column_letter(j)].width = 6
    for i, (name, cells) in enumerate(report["rows"], start=1):
        r = head + i
        vals = [i, name, *cells]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v if v != "" else None)
            c.border = _BORDER
            c.font = _BOLD if v == "M" else _BASE
            c.alignment = Alignment(horizontal="center" if col != 2 else "left")
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 30
    if not report["rows"]:
        ws.cell(row=head + 1, column=2, value="No ranks have an active certificate mapping in this category.").font = _BASE
    ws.freeze_panes = ws.cell(row=head + 1, column=3)
    return wb


def _render_employees(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee_HSE_Matrix"
    certs = report["certs"]
    last_col = 4 + len(certs)
    _logo(ws, "A1")
    ws.cell(row=3, column=4, value="HSE TRAINING MATRIX").font = Font(name="Arial", size=11, bold=True)
    ws.cell(row=4, column=4, value=f"From : {report['from']:%d/%m/%Y}   To: {report['to']:%d/%m/%Y}").font = _BOLD
    ws.cell(row=5, column=4, value=f"Rig : {report['rig'].rig_name}").font = _BOLD
    r = 7
    for label, values in (
        ("Candidates requiring training", report["requiring"]),
        ("Candidates trained", report["trained"]),
        ("Sessions held", report["sessions"]),
    ):
        ws.cell(row=r, column=4, value=label).font = _BOLD
        ws.cell(row=r, column=4).border = _BORDER
        for j, v in enumerate(values, start=5):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border, c.alignment = _BASE, _BORDER, Alignment(horizontal="center")
        r += 1
    head = r + 1
    for col, text in enumerate(("SrNo.", "EMP. NO", "EMPLOYEE NAME", "DESIGNATION"), start=1):
        c = ws.cell(row=head, column=col, value=text)
        c.font, c.border = _BOLD, _BORDER
    for j, cert in enumerate(certs, start=5):
        c = ws.cell(row=head, column=j, value=cert.cert_name)
        c.font, c.border, c.alignment = _BOLD, _BORDER, Alignment(horizontal="center", wrap_text=True, vertical="center")
        if cert.cert_training_type == "E":
            c.fill = _GREY
        ws.column_dimensions[get_column_letter(j)].width = 24
    ws.row_dimensions[head].height = 32
    for k, (n, emp_no, name, rank, cells) in enumerate(report["rows"], start=1):
        row = head + k
        for col, v in enumerate((n, emp_no, name, rank, *cells), start=1):
            c = ws.cell(row=row, column=col, value=v if v != "" else None)
            c.border, c.font = _BORDER, _BASE
            c.alignment = Alignment(horizontal="center" if col > 4 or col == 1 else "left")
    for col, w in ((1, 7), (2, 11), (3, 28), (4, 30)):
        ws.column_dimensions[get_column_letter(col)].width = w
    if not report["rows"]:
        ws.cell(row=head + 1, column=3, value="No employees match this rig and date range.").font = _BASE
    ws.freeze_panes = ws.cell(row=head + 1, column=5)
    return wb


def _payload(report):
    """The same report as the Excel file, as JSON for the on-screen preview."""
    certs = [{"name": c.cert_name, "external": c.cert_training_type == "E"} for c in report["certs"]]
    if report["type"] == "Training_Matrix":
        return {
            "type": report["type"],
            "title": f"HSE TRAINING MATRIX ({_category_label(report['category'])})",
            "certs": certs,
            "rows": [{"no": i, "name": name, "cells": cells} for i, (name, cells) in enumerate(report["rows"], start=1)],
        }
    return {
        "type": report["type"],
        "title": "HSE TRAINING MATRIX",
        "subtitle": [f"From : {report['from']:%d/%m/%Y}   To: {report['to']:%d/%m/%Y}", f"Rig : {report['rig'].rig_name}"],
        "certs": certs,
        "summary": [
            {"label": "Candidates requiring training", "values": report["requiring"]},
            {"label": "Candidates trained", "values": report["trained"]},
            {"label": "Sessions held", "values": report["sessions"]},
        ],
        "rows": [{"no": n, "emp_no": emp_no, "name": name, "designation": rank, "cells": cells} for n, emp_no, name, rank, cells in report["rows"]],
    }


class TrainingReportPreviewView(APIView):
    """The report as data, to show on the page. Needs View only — it is not a download."""

    permission_classes = [HasMenuPermission]
    entity_key = ENTITY_KEY
    action = "list"

    def get(self, request):
        try:
            return Response(_payload(build_report(request)))
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=400)


class TrainingReportCategoriesView(APIView):
    """Categories this user may report on."""

    permission_classes = [HasMenuPermission]
    entity_key = ENTITY_KEY
    action = "list"

    def get(self, request):
        qs = MstFsCategory.objects.order_by("fs_category_name")
        allowed = allowed_category_ids(request)
        if allowed is not None:
            qs = qs.filter(pk__in=allowed)
        return Response([{"id": c.pk, "name": c.fs_category_name} for c in qs])


class TrainingReportExportView(APIView):
    permission_classes = [HasMenuPermission]
    entity_key = ENTITY_KEY
    # A plain APIView has no router-assigned .action, so name the right it needs.
    action = "export"

    def get(self, request):
        try:
            report = build_report(request)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=400)
        wb = _render_matrix(report) if report["type"] == "Training_Matrix" else _render_employees(report)
        buf = io.BytesIO()
        wb.save(buf)
        name = f"{TYPES[report['type']]}.xlsx"
        response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{name}"'
        changes = {"report_type": {"old": None, "new": TYPES[report["type"]]}, "category": {"old": None, "new": report["category"].fs_category_name}}
        if report["type"] == "Employees_Matrix":
            changes.update(
                rig={"old": None, "new": report["rig"].rig_name},
                from_date={"old": None, "new": report["from"].isoformat()},
                to_date={"old": None, "new": report["to"].isoformat()},
            )
        _audit.record_action(request, "export", ENTITY_KEY, record_label=f"Training Report — {TYPES[report['type']]}", changes=changes)
        return response

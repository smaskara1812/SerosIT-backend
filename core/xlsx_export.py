from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")


def xlsx_value(v):
    # openpyxl can write Decimal directly in principle, but not every
    # downstream reader agrees — plain float/int is the portable choice
    # for a numeric cell.
    if isinstance(v, Decimal):
        return float(v)
    return v


def build_xlsx_response(filename, sheet_title, header_lines, columns, rows, wide_columns=(), bold_last_row=False):
    """One shared builder for every "table with a bold/grey/bordered header
    row" export in Drilling Details — CSV can't carry borders, bold text,
    or a header fill at all (it's a plain-text format), so these all
    export as real .xlsx instead.

    header_lines: report-title lines shown above the table (first one
      rendered larger/bold as the title); pass [] for no header block.
    columns: column header strings.
    rows: list of row value-lists, already in column order. Decimal values
      are converted automatically; pass plain numbers/strings/None otherwise.
    wide_columns: 0-based column indexes that should get a wider (40-char)
      column, e.g. a free-text remarks column — everything else is sized
      to its header text.
    bold_last_row: bold + grey-fill the final data row (a Totals row),
      matching how the on-screen table's own Totals row is styled.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel's own sheet-name length limit

    row_num = 1
    for i, line in enumerate(header_lines):
        cell = ws.cell(row=row_num, column=1, value=line)
        if i == 0:
            cell.font = Font(bold=True, size=13)
        row_num += 1
    if header_lines:
        row_num += 1  # blank spacer row before the table

    header_row = row_num
    for col, header in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")

    data_row = header_row
    for row_index, row_values in enumerate(rows):
        data_row += 1
        is_totals_row = bold_last_row and row_index == len(rows) - 1
        for col, value in enumerate(row_values, start=1):
            cell = ws.cell(row=data_row, column=col, value=xlsx_value(value))
            cell.border = BORDER
            if is_totals_row:
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL

    for i, header in enumerate(columns, start=1):
        width = 40 if (i - 1) in wide_columns else max(10, min(28, len(str(header)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

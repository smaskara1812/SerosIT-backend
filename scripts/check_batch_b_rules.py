"""Checks the report and analytics page rules (Batch B) through the real API.

    cd backend
    .venv/bin/python scripts/check_batch_b_rules.py

These pages only read data, so nothing is written; the script still runs inside a rolled-back
transaction like the others. Exits with status 1 if any check fails.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

import django

django.setup()

import datetime
import io

from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import load_workbook
from rest_framework.test import APIClient

from core.models import MstRig

FAILED = []


def check(name, ok, detail=""):
    print(("  ok    " if ok else "  FAIL  ") + name + (f"  {detail[:200]}" if detail and not ok else ""))
    if not ok:
        FAILED.append(name)


def run():
    client = APIClient(HTTP_HOST="localhost")
    client.force_authenticate(get_user_model().objects.filter(is_superuser=True).first())
    today = datetime.date.today()
    day = lambda n: (today + datetime.timedelta(days=n)).isoformat()
    rig = MstRig.objects.first().pk

    pages = (
        ("Operations Analytics", "/api/drilling/operations-analytics/", {"report_type": "daily"}),
        ("Drilling & Tripping Analysis", "/api/drilling/drilling-tripping-analysis/", {"report_type": "drilling"}),
        ("Drilling Daily Data", "/api/drilling/drilling-daily-data/", {}),
    )
    for label, url, extra in pages:
        print(label)
        q = lambda f, t, **o: {"rig": rig, "from_dt": f, "to_dt": t, **extra, **o}
        r = client.get(url, q(day(-30), day(-1)))
        check("a normal past range is accepted", r.status_code == 200, r.content.decode())
        check("today as To is accepted", client.get(url, q(day(-5), day(0))).status_code == 200)
        check("From equal to To is accepted (a one-day report)", client.get(url, q(day(-3), day(-3))).status_code == 200)
        for name, f, t, text in (
            ("a From date in the future", day(1), day(2), "From date can't be in the future"),
            ("a To date in the future", day(-5), day(3), "To date can't be in the future"),
            ("To before From", day(-1), day(-10), "To date can't be before the From date"),
        ):
            r = client.get(url, q(f, t))
            check(f"{name} is refused", r.status_code == 400 and text in r.json().get("error", ""), r.content.decode())
        check("a missing rig is refused", client.get(url, {"from_dt": day(-5), "to_dt": day(-1), **extra}).status_code == 400)
        check("a missing date is refused", client.get(url, {"rig": rig, "from_dt": day(-5), **extra}).status_code == 400)
        check("a bad date is refused", client.get(url, q("not-a-date", day(-1))).status_code == 400)
        r = client.get(url + "export/", q(day(1), day(2)))
        wb = load_workbook(io.BytesIO(b"".join(r.streaming_content) if getattr(r, "streaming", False) else r.content))
        cells = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value]
        check("the Excel export gives the message instead of data for a future date", any("future" in c for c in cells), str(cells[:4]))

    print("HSE Drill Report")
    url = "/api/qhse/hse-drill-report/export/"
    rigs = str(rig)
    this = f"{today.year}-{today.month:02d}"
    prev = f"{today.year - 1}-{today.month:02d}"
    nxt_year, nxt_month = (today.year + 1, 1) if today.month == 12 else (today.year, today.month + 1)
    nxt = f"{nxt_year}-{nxt_month:02d}"
    ok = client.get(url, {"period": "DATE_RANGE", "month": prev, "to_month": this, "rigs": rigs})
    check("a range ending this month is accepted", ok.status_code == 200, ok.content[:150].decode(errors="ignore"))
    r = client.get(url, {"period": "DATE_RANGE", "month": prev, "to_month": nxt, "rigs": rigs})
    check("a To month after the current month is refused", r.status_code == 400 and "after the current month" in r.json().get("detail", ""), r.content.decode(errors="ignore"))
    r = client.get(url, {"period": "DATE_RANGE", "month": this, "to_month": prev, "rigs": rigs})
    check("To before From is refused", r.status_code == 400 and "before" in r.json().get("detail", ""), r.content.decode(errors="ignore"))
    check("a Monthly report still works", client.get(url, {"period": "MONTHLY", "month": prev, "rigs": rigs}).status_code == 200)
    check("a missing period is refused", client.get(url, {"rigs": rigs}).status_code == 400)


if __name__ == "__main__":
    with transaction.atomic():
        run()
        transaction.set_rollback(True)
    print()
    print("FAILED: " + ", ".join(FAILED) if FAILED else "All checks passed. Nothing was saved.")
    sys.exit(1 if FAILED else 0)

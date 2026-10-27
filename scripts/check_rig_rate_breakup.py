"""Checks the Rig Rate Break-Up Sheet mail (alert 233): the Excel against the legacy sample file (every cell), the period (month to date, the whole previous
month on the 1st), which rigs get a sheet, the sheet contents, the mail, the cycle (Next Mail and Start Time), and the Report Mails endpoints. A fake
mail server is used and everything in the database is rolled back.

    cd backend
    .venv/bin/python scripts/check_rig_rate_breakup.py
"""
import datetime, io, os, shutil, sys, tempfile, django
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TMP = tempfile.mkdtemp(prefix="rig_rate_check_")
os.environ["LOG_DIR"] = os.path.join(TMP, "logs")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings"); django.setup()
from django.conf import settings
from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import load_workbook
from rest_framework.test import APIClient
from core import mailer, rig_rate_breakup as rr
from core.models import AlertSenderDtl, MailAlertDtl, MstCostCentre, MstRig, NotificationTrigger, ProjectContract, ProjectContractDtl, SchedItem, UserPermission, UserProfile
from core.permissions import invalidate_user_access

SAMPLE = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "Rig Rate Break-Up", "Rig_Rate_Break_Up_Sheet_15_10_2026_114036.xlsx")
sent = []


def fake_deliver(subject, html, text, to, cc, bcc, from_name, smtp, inline_images, attachments=None):
    sent.append({"subject": subject, "html": html, "to": to, "cc": cc, "bcc": bcc, "from_name": from_name, "host": smtp["host"], "attachments": attachments or []})


mailer._deliver = fake_deliver
settings.EMAIL_HOST = "smtp.default.test"
settings.BIRTHDAY_EMAIL_HOST = "smtp.env.test"
settings.SCHEDULER_ERROR_MAIL_TO = ["it@example.test"]


class Rollback(Exception):
    pass


fails = []


def ok(name, cond, detail=""):
    print(("ok   " if cond else "FAIL ") + name + (f"  {detail}" if not cond else ""))
    if not cond:
        fails.append(name)


def mk(n, grants):
    U = get_user_model(); u = U.objects.create_user(username=n, password="x")
    uid = (UserProfile.objects.order_by("-user_id").first().user_id or 0) + 1
    UserProfile.objects.create(user_id=uid, user_login_id=n, is_app_admin=False)
    for key, flags in grants.items():
        UserPermission.objects.create(user_id=uid, menu_key=key, can_view="view" in flags, can_add="add" in flags, can_edit="edit" in flags, can_delete="delete" in flags, can_export="export" in flags)
    invalidate_user_access(n)
    c = APIClient(HTTP_HOST="localhost"); c.force_authenticate(u); return c


D, DT = datetime.date, datetime.datetime
ALERT = 233
JOB = rr.JOB
NOW = DT(2026, 10, 27, 11, 0, 5)


def alert(**kw):
    MailAlertDtl.objects.filter(pk=ALERT).update(**kw)
    return MailAlertDtl.objects.get(pk=ALERT)


def names(today, window=30):
    return [r.rig_name for r in rr.active_rigs(today, window)]


try:
    with transaction.atomic():
        SchedItem.objects.filter(job_key=rr.JOB_KEY).delete()
        NotificationTrigger.objects.filter(entity_key=rr.ENTITY_KEY, action=rr.ACTION).delete()
        AlertSenderDtl.objects.filter(alert_id=ALERT).delete()
        a233 = alert(alert_window=30, alert_freq=15, next_mail_dt=DT(2026, 10, 16, 11, 19, 13), start_time=11, alert_active="Y", mail_subject="Seros-IT Auto Report - Rig Rate Break-Up Sheet")

        print("A. against the legacy sample (15 Oct 2026)")
        today = D(2026, 10, 15)
        sheets, (start, end) = rr.collect(today, 30)
        ok("the same three rigs, in the same order", [s["rig"] for s in sheets] == ["EE01", "SK02", "Axom Rhino"], str([s["rig"] for s in sheets]))
        new = load_workbook(io.BytesIO(rr.build_workbook(sheets, today)), rich_text=True)
        old = load_workbook(SAMPLE, rich_text=True)
        ok("one sheet per rig, named after the rig (the legacy sheets were Sheet1, Sheet2, Sheet3)", new.sheetnames == ["EE01", "SK02", "Axom Rhino"] and len(old.sheetnames) == 3)
        norm = lambda v: "" if v is None else str(v).strip()

        def cell_diffs(ns, os_):
            out = []
            for r in range(1, os_.max_row + 1):
                for c in range(1, 12):
                    o, n = norm(os_.cell(r, c).value).replace("SIS Report", "Seros-IT Report"), norm(ns.cell(r, c).value)
                    try:
                        if o != n and float(o or 0) == float(n or 0):
                            continue
                    except ValueError:
                        pass
                    if o != n:
                        out.append((r, c))
            return out
        d = [cell_diffs(ns, os_) for ns, os_ in zip(new.worksheets, old.worksheets)]
        ok("SK02 and Axom Rhino: every cell is the same as the legacy file (apart from SIS -> Seros-IT)", d[1] == [] and d[2] == [], str(d[1][:5] + d[2][:5]))
        ok("EE01 differs only on 14 Oct (entered after our 14 Oct data dump) and in the totals that include it", d[0] and all(r in (24, 26) for r, _ in d[0]) and new.worksheets[0]["C26"].value == 295.25, str(d[0]))
        ws = new.worksheets[0]
        ok("same size and widths as the legacy sheet", ws.max_row == old.worksheets[0].max_row and {k: v.width for k, v in ws.column_dimensions.items()} == {k: v.width for k, v in old.worksheets[0].column_dimensions.items()})
        ok("the header: company, title, report date (Seros-IT), rig, locations, area, month", ws["A1"].value == "Seros Drilling Private Limited" and ws["A2"].value == "Rig Rate Break-Up Sheet" and "Seros-IT Report Date: 15-10-2026" == str(ws["A3"].value)
           and str(ws["A4"].value) == "Rig:  EE01" and str(ws["A5"].value) == "Locations : BKH" and str(ws["A6"].value) == "Area: Duliajan" and ws["A7"].value == "October-2026")
        ok("the company header is blue, bold, 15 pt and centred", ws["A1"].font.sz == 15 and ws["A1"].font.b and ws["A1"].font.color.rgb.endswith("0000FF") and ws["A1"].alignment.horizontal == "center")
        ok("the column heads are the rate codes", [c.value for c in ws[10]] == ["Date", "Well", "R1", "R2", "R3", "R0", "FM", "ILM", "IPM - Lateral", "Total", "Operations Summary"])
        ok("hours are numbers, a day with no report is a blank row, and Total Hours closes the sheet", isinstance(ws["C17"].value, float) and ws["C17"].value == 7.25 and ws["C25"].value is None and ws["A26"].value == "Total Hours")
        ok("the Operations Summary is the day's remark", ws["K11"].value.startswith("RIH 12.1/4"))

        print("B. the period")
        ok("month to date", rr.period(D(2026, 10, 15)) == (D(2026, 10, 1), D(2026, 10, 15)))
        ok("on the 1st: the whole previous month", rr.period(D(2026, 11, 1)) == (D(2026, 10, 1), D(2026, 10, 31)))
        ok("on 1 January: December of last year", rr.period(D(2027, 1, 1)) == (D(2026, 12, 1), D(2026, 12, 31)))
        ok("and after a leap February", rr.period(D(2028, 3, 1)) == (D(2028, 2, 1), D(2028, 2, 29)))
        ok("the month label is English and does not depend on the server's language", rr.month_label(D(2026, 10, 31)) == "October-2026" and rr.month_label(D(2026, 3, 1)) == "March-2026")
        s1, (a1, b1) = rr.collect(D(2026, 11, 1), 30)
        ok("a sheet made on the 1st covers every day of the previous month, with that month's name", s1 and len(s1[0]["rows"]) == 31 and s1[0]["month"] == "October-2026" and s1[0]["rows"][0]["date"] == D(2026, 10, 1) and s1[0]["rows"][-1]["date"] == D(2026, 10, 31))
        s2, _ = rr.collect(today, 30)
        ok("a mid-month sheet has one row per day up to today, today's included", len(s2[0]["rows"]) == 15 and s2[0]["rows"][-1]["date"] == today)

        print("C. which rigs get a sheet")
        ok("the window is how far back a rig must have a daily report: a window of 0 needs one yesterday (none here)", names(today, 0) == [])
        ok("a window of 30 finds the three", names(today, 30) == ["EE01", "SK02", "Axom Rhino"])
        ok("far in the future no rig has a recent report", names(D(2027, 6, 1), 30) == [])
        rig = MstRig.objects.get(rig_name="EE01")
        MstCostCentre.objects.filter(rig_id=rig.rig_id).update(cost_centre_active="N")
        ok("a rig whose cost centre is not active is left out", names(today) == ["SK02", "Axom Rhino"])
        MstCostCentre.objects.filter(rig_id=rig.rig_id).update(cost_centre_active="Y")
        ProjectContractDtl.objects.filter(rig_id=rig.rig_id, rig_active_to__isnull=True).update(rig_active_to=D(2026, 9, 1))
        ok("a rig with no active contract line is left out", names(today) == ["SK02", "Axom Rhino"])
        ProjectContractDtl.objects.filter(rig_id=rig.rig_id, rig_active_to=D(2026, 9, 1)).update(rig_active_to=None)
        ok("and is back when the line is active again", names(today) == ["EE01", "SK02", "Axom Rhino"])
        ln = ProjectContractDtl.objects.filter(rig_id=rig.rig_id, rig_active_to__isnull=True).first()
        ProjectContract.objects.filter(pk=ln.contract_id).update(prj_end_dt=D(2026, 9, 30))
        ok("a rig whose contract has ended is left out", "EE01" not in names(today))

        print("D. the mail")
        item_mail = rr.build_mail(a233, DT(2026, 10, 15, 11, 40, 36))
        ok("the subject is the alert's subject", item_mail.subject == "Seros-IT Auto Report - Rig Rate Break-Up Sheet")
        ok("the heading says Seros-IT Auto Report, the sentence has the date, there is no 'S.I.S.'", "Seros-IT Auto Report" in item_mail.html and "Please find attached Rig Rate Break-Up Sheet dated 15-10-2026" in item_mail.html and "S.I.S." not in item_mail.html and "automated Mail" in item_mail.html)
        ok("the file name follows the legacy pattern", item_mail.attachments[0][0] == "Rig_Rate_Break_Up_Sheet_15_10_2026_114036.xlsx" and "spreadsheetml" in item_mail.attachments[0][2])
        empty = rr.build_mail(alert(alert_window=1), DT(2027, 6, 1, 11, 0, 0))
        ok("with no rig the mail says so, with no attachment", empty.empty and not empty.attachments and "No rig has a daily drilling report in the last 1 days" in empty.html, empty.html[:200])
        alert(alert_window=30)
        ok("sheet names are made safe and unique", rr._sheet_name("A/B:C", set()) == "A B C" and rr._sheet_name("X", {"x"}) == "X (2)" and len(rr._sheet_name("n" * 60, set())) <= 31)

        print("E. the cycle: Next Mail date and Start Time")
        ProjectContract.objects.filter(pk=ln.contract_id).update(prj_end_dt=None)
        s = JOB.get_setup(NOW.date())
        ok("with no trigger the page is told what to do", not s["trigger_found"] and s["warnings"])
        did = rr.ensure_setup()
        ok("the setup command finds the alert by its name and creates the trigger, once", NotificationTrigger.objects.filter(entity_key="automation", action="send_rig_rate_breakup_mail", alert_id=ALERT).exists() and did and rr.ensure_setup() == [])
        s = JOB.get_setup(NOW.date())
        ok("the setup shows the cycle from the alert: every 15 days, window 30, start hour 11", s["freq"] == 15 and s["window"] == 30 and s["send_hour"] == 11 and s["next_mail_dt"] == DT(2026, 10, 16, 11, 19, 13) and "Every 15 days" in JOB.schedule_text(NOW.date()))
        ok("the From name is blank (the mail server adds its own) and the default account is used", s["from_name"] == "" and s["smtp"]["sender"] is None)
        alert(alert_window=0)
        ok("a window of 0 is a warning (the legacy report used 30)", any("Window is 0" in w for w in JOB.get_setup(NOW.date())["warnings"]))
        alert(alert_window=30)
        JOB.plan(NOW)
        it = SchedItem.objects.get(job_key=rr.JOB_KEY)
        ok("a Next Mail date that has passed is planned for today, at the Start Time (11:00) entered on the alert", it.send_date == NOW.date() and it.due_at == DT(2026, 10, 27, 11, 0) and "Due 16 Oct 2026" in it.note, str((it.send_date, it.due_at, it.note)))
        n0 = len(sent); JOB.process(DT(2026, 10, 27, 10, 59, 0))
        ok("nothing goes before the hour", len(sent) == n0)
        JOB.process(NOW)
        m = sent[-1]
        ok("at 11:00 it goes out, to the alert's To, Cc and Bcc, through the default account, with the Excel", len(sent) == n0 + 1 and m["host"] == "smtp.default.test" and m["to"] and len(m["attachments"]) == 1 and m["from_name"] == "")
        ok("the recipients are the alert's own rows (DR01, EE01, Axom Rhino in To)", {"ds.dr01@seros.co.in", "ds.ee01@seros.co.in", "rm.axomrhino@seros.co.in"} <= {x.lower() for x in m["to"]}, str(m["to"]))
        a = MailAlertDtl.objects.get(pk=ALERT)
        ok("Next Mail moves on by the frequency from the day it was due: 16 Oct -> 31 Oct, at 11:00", a.next_mail_dt == DT(2026, 10, 31, 11, 0), str(a.next_mail_dt))
        ok("the item says what was sent", "rigs" in SchedItem.objects.get(job_key=rr.JOB_KEY).note)
        alert(next_mail_dt=DT(2026, 11, 3, 14, 0), start_time=14)
        SchedItem.objects.filter(job_key=rr.JOB_KEY).delete()
        JOB.plan(DT(2026, 10, 30, 8, 0, 0))
        it = SchedItem.objects.get(job_key=rr.JOB_KEY)
        ok("a date set on the form is the date, and the time on the form is the time: planned for 3 Nov at 14:00", it.send_date == D(2026, 11, 3) and it.due_at == DT(2026, 11, 3, 14, 0))
        n1 = len(sent); JOB.process(DT(2026, 11, 3, 13, 59, 0)); JOB.process(DT(2026, 11, 3, 14, 0, 5))
        ok("and it goes at 14:00 on that day, not before", len(sent) == n1 + 1 and MailAlertDtl.objects.get(pk=ALERT).next_mail_dt == DT(2026, 11, 18, 14, 0), str(MailAlertDtl.objects.get(pk=ALERT).next_mail_dt))

        print("F. the pages")
        full = mk("zzrr1", {"admin.report_mails": ("view", "edit", "export"), "admin.scheduled_mails": ("view",)})
        alert(next_mail_dt=DT(2026, 10, 16, 11, 19, 13), start_time=11)
        cards = full.get("/api/admin/report-mails/").json()["results"]
        ok("Report Mails lists both reports", {c["key"] for c in cards} == {"haz_id_register", "rig_rate_breakup"})
        c = next(x for x in cards if x["key"] == "rig_rate_breakup")
        ok("the card has the setup, the schedule and the recipients", c["name"] == "Rig Rate Break-Up Sheet" and c["setup"]["alert"]["id"] == ALERT and "Every 15 days" in c["schedule"] and c["preview"]["counts"]["to"] >= 3 and c["overdue"])
        r = full.get("/api/admin/report-mails/rig_rate_breakup/download/")
        ok("the Excel can be downloaded (Export)", r.status_code == 200 and "spreadsheetml" in r["Content-Type"] and r["Content-Disposition"].startswith('attachment; filename="Rig_Rate_Break_Up_Sheet_') and load_workbook(io.BytesIO(r.content)).sheetnames[0])
        sent.clear()
        r = full.post("/api/admin/report-mails/rig_rate_breakup/test-send/", {"email": "me@example.test"}, format="json")
        ok("a test goes to one address, marked [Test], with the Excel, and does not move Next Mail", r.status_code == 200 and len(sent) == 1 and sent[0]["to"] == ["me@example.test"] and sent[0]["subject"].startswith("[Test] ") and sent[0]["attachments"] and MailAlertDtl.objects.get(pk=ALERT).next_mail_dt == DT(2026, 10, 16, 11, 19, 13), str(r.content[:150]))
        sm = full.get("/api/admin/scheduled-mails/").json()
        ok("Scheduled Mails lists the four automations", {j["key"] for j in sm["jobs"]} == {"birthday", "holiday", "haz_id_register", "rig_rate_breakup"})
        raise Rollback()
except Rollback:
    pass
finally:
    shutil.rmtree(TMP, ignore_errors=True)

print()
print("FAILED: " + ", ".join(fails) if fails else "all checks passed")
sys.exit(1 if fails else 0)
